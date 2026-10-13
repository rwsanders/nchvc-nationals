"""Pull the NCHVC Nationals bracket workbooks from Google Sheets and normalize them to JSON.

Usage:  python nchvc_scrape.py [--cache DIR] [--offline] [--out data.json]
"""
import argparse, datetime as dt, io, json, os, re, sys, time, urllib.request, zipfile
from concurrent.futures import ThreadPoolExecutor

import openpyxl
from openpyxl.utils import get_column_letter

INDEX_ID = "1Uq8uFxIn-y3YwMAzIwCft1l4zgfFA0E7rKGuIBQvo5U"
EXPORT = "https://docs.google.com/spreadsheets/d/{}/export?format=xlsx"
SHEET_URL = "https://docs.google.com/spreadsheets/d/{}/edit"
DAYS = {"Thursday", "Friday", "Saturday", "Sunday", "Wednesday"}


# ---------------------------------------------------------------- download
def fetch(sheet_id, cache_dir, offline=False, tries=4):
    path = os.path.join(cache_dir, sheet_id + ".xlsx")
    if offline:
        return path
    last = None
    for i in range(tries):
        try:
            req = urllib.request.Request(EXPORT.format(sheet_id), headers={"User-Agent": "Mozilla/5.0"})
            data = urllib.request.urlopen(req, timeout=120).read()
            zipfile.ZipFile(io.BytesIO(data)).testzip()  # reject truncated downloads
            with open(path, "wb") as f:
                f.write(data)
            return path
        except Exception as e:  # noqa: BLE001
            last = e
            time.sleep(2 * (i + 1))
    if os.path.exists(path):
        print(f"  ! {sheet_id}: download failed ({last}); using cached copy", file=sys.stderr)
        return path
    raise RuntimeError(f"could not download {sheet_id}: {last}")


# ---------------------------------------------------------------- helpers
def s(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def fmt_time(v):
    if isinstance(v, dt.datetime):
        v = v.time()
    if isinstance(v, dt.time):
        h = v.hour % 12 or 12
        return f"{h}:{v.minute:02d} {'AM' if v.hour < 12 else 'PM'}"
    if isinstance(v, float) and 0 <= v < 1:  # fractional day
        mins = round(v * 24 * 60)
        return fmt_time(dt.time(mins // 60 % 24, mins % 60))
    t = s(v)
    m = re.fullmatch(r"(\d{1,2}):(\d{2})(?::\d{2})?", t)
    if m:
        return fmt_time(dt.time(int(m[1]) % 24, int(m[2])))
    return t


def num(v):
    try:
        return int(float(v))
    except (TypeError, ValueError):
        return None


class Grid:
    def __init__(self, ws):
        self.title = ws.title
        self.v = {}
        for row in ws.iter_rows():
            for c in row:
                if c.value not in (None, ""):
                    self.v[(c.row, c.column)] = c.value

    def __call__(self, r, c):
        return self.v.get((r, c))

    def t(self, r, c):
        return s(self.v.get((r, c)))

    def find(self, pred):
        return sorted((rc for rc, val in self.v.items() if pred(s(val))), key=lambda rc: rc)


def parse_sets(text):
    """'25-8, 25-12' -> [[25,8],[25,12]]"""
    return [[int(a), int(b)] for a, b in re.findall(r"(\d+)\s*-\s*(\d+)", text or "")]


# ---------------------------------------------------------------- pools
def parse_pools(g, prefix=""):
    pools = []
    for (r, c) in g.find(lambda x: re.fullmatch(r"(Iron )?Pool [A-Z]+ Schedule.*", x) is not None):
        name = re.sub(r" Schedule.*", "", g.t(r, c))
        best_of = 3 if "Best of 3" in g.t(r, c) else 2
        matches, rr = [], r + 1
        while rr < r + 12:
            label = g.t(rr, c)
            if label.startswith("Match #"):
                teams = [g.t(rr, cc) for cc in range(c + 4, c + 12) if g.t(rr, cc) and g.t(rr, cc) != "vs"]
                matches.append({"n": num(label[7:]), "court": g.t(rr, c + 1), "time": fmt_time(g(rr, c + 2)),
                                "a": teams[0] if teams else "", "b": teams[1] if len(teams) > 1 else ""})
            elif label.startswith("Seed"):
                break
            rr += 1
        # score grid: team rows have numeric seed in col c, name in c+1; groups of 3 cols from c+3
        seed_row = rr
        teams, scores = [], {}
        rr = seed_row + 2
        while rr < seed_row + 30:
            seed = num(g(rr, c))
            team = g.t(rr, c + 1)
            if seed is None:
                if g.t(rr, c).endswith("Results") or g.t(rr, c).startswith("Pool"):
                    break
                rr += 1
                continue
            row = {"seed": seed, "team": team, "diff": num(g(rr, c + 2)) or 0, "_r": rr}
            teams.append(row)
            for gc in range(c + 3, c + 3 + 3 * 6, 3):
                me, opp = g.t(rr, gc + 1), g.t(rr, gc + 2)
                if not opp or g.t(rr + 1, gc) != "Set 1":
                    continue
                sets = []
                for k in range(best_of):
                    a, b = num(g(rr + 1 + k, gc + 1)), num(g(rr + 1 + k, gc + 2))
                    if a is not None and b is not None and (a or b):
                        sets.append([a, b])
                scores[(team, opp)] = sets
            rr += best_of + 1
        # Win / Loss header
        hdr = [(cc, g.t(seed_row + 1, cc)) for cc in range(c, c + 24)]
        wcol = next((cc for cc, t in hdr if t == "Win"), None)
        for t in teams:
            r0 = t.pop("_r")
            if wcol:
                t["w"], t["l"] = num(g(r0, wcol)) or 0, num(g(r0, wcol + 1)) or 0
        for m in matches:
            sets = scores.get((m["a"], m["b"]))
            if sets is None and (m["b"], m["a"]) in scores:
                sets = [[b, a] for a, b in scores[(m["b"], m["a"])]]
            m["sets"] = sets or []
            m["status"] = match_status(m["sets"], best_of)
            if m["status"] == "final":
                aw = sum(1 for a, b in m["sets"] if a > b)
                bw = len(m["sets"]) - aw
                m["winner"] = m["a"] if aw > bw else m["b"] if bw > aw else None  # pool sets can split 1-1
        # results table
        standings = []
        res = g.find(lambda x: x == f"{name} Results")
        if res:
            rr, rc = res[0]
            hdr = {g.t(rr + 1, cc): cc for cc in range(rc, rc + 12) if g.t(rr + 1, cc)}
            tcol, acol = hdr.get("Team", rc + 2), hdr.get("Advancement", rc + 5)
            done = any(g.t(rr, cc) == "complete" or g.t(rr + 2, cc) == "complete" for cc in range(rc, rc + 24))
            k = rr + 2
            while g.t(k, rc) and re.fullmatch(r"\d+(st|nd|rd|th)", g.t(k, rc)):
                standings.append({"rank": g.t(k, rc), "team": g.t(k, tcol), "adv": g.t(k, acol)})
                k += 1
        pools.append({"name": f"{prefix}{name}", "bestOf": best_of, "court": matches[0]["court"] if matches else "",
                      "teams": teams, "matches": matches, "standings": standings})
    return pools


def match_status(sets, best_of):
    if not sets:
        return "upcoming"
    if best_of == 2:  # pool play: always 2 sets
        return "final" if len(sets) >= 2 else "live"
    need = best_of // 2 + 1
    a = sum(1 for x, y in sets if x > y)
    return "final" if max(a, len(sets) - a) >= need else "live"


# ---------------------------------------------------------------- brackets
def parse_bracket_tab(g):
    blocks = []
    for (r, c) in g.find(lambda x: x in DAYS):
        if not (g(r + 1, c) is not None and g.t(r + 2, c) and g.t(r + 3, c)):
            continue
        name = g.t(r + 3, c)
        best = re.search(r"Best of (\d)", g.t(r - 1, c) + " " + name)
        def slot(rows, default):
            # the seed label (col c-1) marks the slot row; it is filled in even before teams are known
            sr = next((x for x in rows if g.t(x, c - 1)), default)
            score = next((g.t(x, c) for x in (sr + 1, sr + 2) if parse_sets(g.t(x, c))), "")
            return {"seed": g.t(sr, c - 1), "team": g.t(sr, c), "score": score}
        top, bot = slot((r - 3, r - 4, r - 5), r - 3), slot((r + 5, r + 4, r + 6), r + 5)
        loser_to = loser_name = ""
        near = [(abs(rr - r), rr, cc) for (rr, cc), v in g.v.items()
                if s(v) == "Losing team to" and c + 2 <= cc <= c + 4 and abs(rr - r) <= 6]
        if near:
            _, rr, cc = min(near)
            loser_to, loser_name = g.t(rr + 1, cc), g.t(rr + 2, cc)
            if loser_name.endswith("Below"):  # e.g. "7th Place Game Below"
                loser_to, loser_name = loser_name.replace(" Below", ""), ""
            if loser_to.endswith("Below"):
                loser_to = loser_to.replace(" Below", "")
        blocks.append({"row": r, "col": c, "day": g.t(r, c), "time": fmt_time(g(r + 1, c)), "court": g.t(r + 2, c),
                       "name": name, "bestOf": int(best[1]) if best else 3, "top": top, "bottom": bot,
                       "loserTo": loser_to, "loser": loser_name})
    # winners: a team that shows up in a later-round slot, else the reported loser, else set count
    later = {}
    for b in blocks:
        for sl in (b["top"], b["bottom"]):
            if sl["team"]:
                later.setdefault(sl["team"], []).append(b["col"])
    for b in blocks:
        t, u = b["top"]["team"], b["bottom"]["team"]
        sets = parse_sets(b["top"]["score"]) or parse_sets(b["bottom"]["score"])
        b["sets"] = sets
        w = None
        if t and u:
            if b["loser"] in (t, u):
                w = u if b["loser"] == t else t
            elif any(cc > b["col"] for cc in later.get(t, [])) and not any(cc > b["col"] for cc in later.get(u, [])):
                w = t
            elif any(cc > b["col"] for cc in later.get(u, [])) and not any(cc > b["col"] for cc in later.get(t, [])):
                w = u
            elif sets and match_status(sets, b["bestOf"]) == "final":
                # score is written under the winner, winner's points first
                w = t if b["top"]["score"] else u
        b["winner"] = w
        b["status"] = "final" if w else ("live" if sets else "upcoming")
    # cluster into sub-brackets by overlapping vertical span
    spans = sorted(blocks, key=lambda b: b["row"])
    clusters = []
    for b in spans:
        lo, hi = b["row"] - 3, b["row"] + 6
        for cl in clusters:
            if lo <= cl["hi"] and hi >= cl["lo"]:
                cl["m"].append(b); cl["lo"] = min(cl["lo"], lo); cl["hi"] = max(cl["hi"], hi)
                break
        else:
            clusters.append({"lo": lo, "hi": hi, "m": [b]})
    # merge clusters that now overlap after growth
    merged = True
    while merged:
        merged = False
        for i in range(len(clusters)):
            for j in range(i + 1, len(clusters)):
                a, b = clusters[i], clusters[j]
                if a["lo"] <= b["hi"] and a["hi"] >= b["lo"]:
                    a["m"] += b["m"]; a["lo"] = min(a["lo"], b["lo"]); a["hi"] = max(a["hi"], b["hi"])
                    del clusters[j]; merged = True; break
            if merged:
                break
    out = []
    for cl in sorted(clusters, key=lambda c: c["lo"]):
        ms = sorted(cl["m"], key=lambda b: (b["col"], b["row"]))
        cols = sorted({b["col"] for b in ms})
        for b in ms:
            b["round"] = cols.index(b["col"])
        seeds = [b[k]["seed"] for b in ms for k in ("top", "bottom")]
        pref = [m[1] for x in seeds if (m := re.match(r"(\d+A) #", x))]
        label = max(set(pref), key=pref.count) if pref else ""
        out.append({"label": label, "rows": [cl["lo"], cl["hi"]], "rounds": len(cols), "matches": ms})
    # result tables / champions
    results = []
    for (r, c) in g.find(lambda x: x.endswith(" Results") and "Pool" not in x):
        rows, k = [], r + 2
        hdr = {g.t(r + 1, cc): cc for cc in range(c, c + 6) if g.t(r + 1, cc)}
        tcol, acol = hdr.get("Teams", c + 2), hdr.get("Advancement", c + 1)
        while re.fullmatch(r"\d+(st|nd|rd|th)", g.t(k, c)):
            rows.append({"rank": g.t(k, c), "adv": g.t(k, acol), "team": g.t(k, tcol)})
            k += 1
        results.append({"title": g.t(r, c), "rows": rows})
    return {"clusters": out, "results": results}


# ---------------------------------------------------------------- reference table (hidden "Ref" tab)
def parse_ref(g):
    rows = []
    r = 6
    while r < 400:
        pool, seed, team = g.t(r, 7), g.t(r, 8), g.t(r, 9)
        if pool or team:
            rows.append({"pool": pool, "seed": seed, "team": team, "slot": g.t(r, 10), "adv": g.t(r, 11)})
        elif r > 40:
            break
        r += 1
    return rows


# ---------------------------------------------------------------- index
def parse_index(path):
    wb = openpyxl.load_workbook(path)
    ws = wb["Nationals Index"]
    links, info, gender = [], {}, None
    for row in ws.iter_rows():
        for c in row:
            v = s(c.value)
            if v in ("Girls", "Boys"):
                gender = v
            href = c.hyperlink.target if c.hyperlink else ""
            m = re.search(r"/d/([\w-]+)", href or "")
            if m and "spreadsheets" in href:
                links.append({"label": re.sub(r"\s+", " ", v), "id": m[1], "gender": gender})
            elif href:
                info[v] = href
            elif v.startswith("REPORT SCORES"):
                info["_scores"] = v
            elif "TEAMS" in v:
                info["_teams"] = v
            elif v.startswith("NATIONALS"):
                info["_title"] = v
            elif v.startswith("Please Note"):
                info["_note"] = v
    return links, info


def parse_workbook(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    out = {"pools": [], "brackets": [], "ref": [], "notices": []}
    for ws in wb.worksheets:
        title = ws.title.strip()
        g = Grid(ws)
        if title == "Ref":
            out["ref"] = parse_ref(g)
        elif title == "Home":
            for (r, c) in g.find(lambda x: x.startswith("Important Notice")):
                txt = g.t(r + 1, c)
                if txt:
                    out["notices"].append(txt.lstrip("● ").strip())
        elif "Pools" in title:
            out["pools"] += parse_pools(g)
        else:
            b = parse_bracket_tab(g)
            if b["clusters"]:
                out["brackets"].append({"tab": re.sub(r"^D\d\s+", "", title), **b})
    return out


AGE_RE = re.compile(r"([GB])(\d+u|JV)", re.I)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--cache", default=os.path.join(os.path.dirname(os.path.abspath(__file__)), ".cache"))
    ap.add_argument("--out", default=os.path.join(os.path.dirname(os.path.abspath(__file__)), "data.json"))
    ap.add_argument("--offline", action="store_true", help="reuse cached xlsx files")
    a = ap.parse_args()
    os.makedirs(a.cache, exist_ok=True)

    t0 = time.time()
    links, info = parse_index(fetch(INDEX_ID, a.cache, a.offline))
    print(f"index: {len(links)} division sheets")
    with ThreadPoolExecutor(6) as ex:
        paths = list(ex.map(lambda l: fetch(l["id"], a.cache, a.offline), links))
    sheets = []
    for l, p in zip(links, paths):
        d = parse_workbook(p)
        m = AGE_RE.search(l["label"])
        age = (m[1].upper() + m[2]) if m else l["label"]
        dm = re.search(r"Division (\d)(?:\s*&\s*(\d))?", l["label"])
        sheets.append({"id": l["id"], "label": l["label"], "gender": l["gender"], "age": age,
                       "divs": [x for x in (dm.groups() if dm else ()) if x],
                       "kind": "pools" if "POOLS" in l["label"].upper() else "bracket",
                       "url": SHEET_URL.format(l["id"]), **{k: v for k, v in d.items() if k != "ref"}})
        nm = sum(len(c["matches"]) for b in d["brackets"] for c in b["clusters"])
        print(f"  {l['label']:<28} pools={len(d['pools']):2} bracket-matches={nm:3} ref={len(d['ref'])}")
    data = {"generated": dt.datetime.now(dt.timezone.utc).isoformat(timespec="seconds"),
            "source": f"https://docs.google.com/spreadsheets/d/{INDEX_ID}/edit", "info": info, "sheets": sheets}
    with open(a.out, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, separators=(",", ":"))
    print(f"wrote {a.out} ({os.path.getsize(a.out)//1024} KB) in {time.time()-t0:.0f}s")


if __name__ == "__main__":
    main()
